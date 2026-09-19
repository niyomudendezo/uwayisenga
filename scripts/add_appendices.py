"""Create a UoK-formatted final-project report with supporting appendices."""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path('/home/edy/Videos/Joselyne.docx')
OUTPUT = Path('/home/edy/Videos/Joselyne_UoK_Final_Project.docx')
PROJECT = Path('/opt/lampp/htdocs/agrukrwanda')

def configure_document(doc):
    for section in doc.sections:
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
        section.left_margin, section.right_margin = Inches(1.5), Inches(1.5)
        section.top_margin, section.bottom_margin = Inches(1), Inches(1)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Times New Roman', Pt(12)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    for p in doc.paragraphs:
        if p.style.name in ('Normal', 'List Paragraph'):
            p.paragraph_format.line_spacing, p.paragraph_format.space_after = 1.5, Pt(0)
            if p.text.strip(): p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    update = OxmlElement('w:updateFields')
    update.set(qn('w:val'), 'true')
    doc.settings.element.append(update)

def body(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing, p.paragraph_format.space_after = 1.5, Pt(6)
    p.add_run(text)

def bullet(doc, text):
    p = doc.add_paragraph(style='List Paragraph')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Inches(0.3), Inches(-0.2)
    p.add_run('•  ' + text)

def add_code(doc, title, source, first_line, last_line):
    doc.add_heading(title, level=2)
    note = doc.add_paragraph()
    note.add_run(f'Source: {source}; lines {first_line}–{last_line}.').italic = True
    lines = (PROJECT / source).read_text(encoding='utf-8').splitlines()[first_line - 1:last_line]
    p = doc.add_paragraph()
    p.paragraph_format.left_indent, p.paragraph_format.right_indent = Inches(0.08), Inches(0.05)
    p.paragraph_format.line_spacing, p.paragraph_format.space_after = 1, Pt(8)
    run = p.add_run('\n'.join(f'{i:>3}  {line}' for i, line in enumerate(lines, first_line)))
    run.font.name, run.font.size = 'Courier New', Pt(7.5)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Courier New')

doc = Document(SOURCE)
configure_document(doc)
doc.add_page_break()
heading = doc.add_heading('APPENDICES', level=1)
heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading('APPENDIX I: DATA-COLLECTION INSTRUMENT', level=1)
body(doc, 'The following semi-structured interview guide supports the data-collection methods described in Section 3.3. It was designed for farmers, cooperative managers, buyers, and agricultural-sector administrators. Participation should be voluntary. No respondent should be required to provide a name, national identification number, password, or other unnecessary personal information.')
doc.add_heading('Introduction to the Respondent', level=2)
body(doc, 'Dear respondent, this interview is conducted for the AgruKrwanda final-year project at the University of Kigali. Its purpose is to understand challenges in cooperative record management, market information, inventory management, and buyer linkage. The information will be used only for academic and system-design purposes. You may decline to answer any question or stop the interview at any time.')
for title, questions in [
    ('Section A: Respondent Profile', ['What is your role: farmer, cooperative manager, buyer, or agricultural-sector administrator?', 'For how many years have you worked in this role or participated in agricultural marketing?', 'Which district do you operate in?', 'What crops do you produce, manage, purchase, or monitor most often?']),
    ('Section B: Current Practices and Challenges', ['How do you currently obtain or share market-price information?', 'What difficulties do you experience when recording harvests, stock, sales, or orders?', 'How are buyers and sellers currently connected, and what challenges arise in this process?', 'What information would help you make better production, pricing, or purchasing decisions?', 'What risks or concerns would you have when using a web-based agricultural system?']),
    ('Section C: System Requirements and Evaluation', ['Which functions would be most useful: price information, harvest recording, inventory management, buyer orders, reports, or forecasts? Please explain.', 'What language, training, device access, or connectivity support would you need to use the system?', 'After viewing or using a prototype, what was easy to use and what should be improved?', 'Would you be willing to use AgruKrwanda in your regular work? Why or why not?']),
]:
    doc.add_heading(title, level=2)
    for question in questions: bullet(doc, question)

doc.add_heading('APPENDIX II: SELECTED SOURCE-CODE LISTINGS', level=1)
body(doc, 'This appendix contains selected code written for the AgruKrwanda project and directly supports the implementation claims in Chapter Four. The full source code is submitted electronically as a separate archive because reproducing the complete codebase in print would make the report unnecessarily long. Third-party framework files and external libraries are not reproduced as project source code.')
add_code(doc, 'II.1 Application Entry Point and Security Headers', 'public/index.php', 1, 21)
add_code(doc, 'II.2 Session, Role, and CSRF Functions', 'app/helpers/Auth.php', 1, 58)
add_code(doc, 'II.3 PDO Database Connection', 'config/database.php', 1, 26)
add_code(doc, 'II.4 Selected Role-Based Routes', 'routes/web.php', 1, 54)
add_code(doc, 'II.5 Local Price-Prediction Workflow', 'app/services/AIPredictionService.php', 25, 91)
add_code(doc, 'II.6 Linear Regression, Demand, and Confidence Logic', 'app/services/AIPredictionService.php', 207, 258)

doc.add_heading('APPENDIX III: ELECTRONIC SOURCE-CODE SUBMISSION', level=1)
body(doc, 'The complete AgruKrwanda source-code archive accompanies this final project report. It contains the application source files, database schema and seed data, route definitions, configuration templates, public assets, and project documentation. The archive is the authoritative full implementation; Appendix II is included only to provide assessors with readable evidence of the principal implementation components.')
for item in ['Archive name: AgruKrwanda_Source_Code_Submission.zip.', 'Technology: PHP 8 or later, MySQL 8 or compatible MariaDB, Apache, HTML, CSS, JavaScript, Bootstrap, and Chart.js.', 'Database installation: create the agrukrwanda database and import database/schema.sql followed by database/seeds.sql.', 'Application entry point: public/index.php. Configure the database connection in config/database.php before deployment.', 'Integrity statement: only project files needed to build, run, and assess the system are included; large demonstration media and generated screenshots are excluded from the code archive.']:
    bullet(doc, item)

doc.save(OUTPUT)
print(OUTPUT)
