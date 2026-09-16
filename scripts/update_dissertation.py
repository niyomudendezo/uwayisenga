from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt, RGBColor, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path('/home/edy/Pictures/Joselyne.docx')
OUTPUT = Path('/home/edy/Pictures/Joselyne_Updated_Chapters_4_5.docx')
SCREENSHOTS = Path('/opt/lampp/htdocs/agrukrwanda/public/assets/images/screenshots')

doc = Document(SOURCE)

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)

def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)

def body(text, bold_lead=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p

def bullet(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(3)
    p.add_run('•  ' + text)
    return p

def heading(text, level):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p

def caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text).bold = True
    return p

def figure(filename, caption_text, width=6.25):
    image_path = SCREENSHOTS / filename
    if not image_path.exists():
        raise RuntimeError(f'Required screenshot was not found: {image_path}')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(image_path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    cap.add_run(caption_text).bold = True
    return p

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    hdr = t.rows[0]
    set_repeat_table_header(hdr)
    for i, value in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, '198754')
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        run = cell.paragraphs[0].add_run(value)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[i].text = str(value)
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(8.5)
    doc.add_paragraph()
    return t

# Locate References and update the earlier chapter-organization statement.
references = next((p for p in doc.paragraphs if p.text.strip().upper() == 'REFERENCES'), None)
if references is None:
    raise RuntimeError('REFERENCES heading was not found in the dissertation.')

for p in doc.paragraphs:
    if p.text.strip().startswith('This report is organized into three chapters.'):
        p.text = ('This report is organized into five chapters. Chapter One presents the general introduction and '
                  'background to the study. Chapter Two reviews related literature on agricultural market systems, '
                  'cooperative management, and AI in agriculture. Chapter Three describes the research methodology '
                  'and system design. Chapter Four presents the implementation, testing results, and discussion of the '
                  'developed AgruKrwanda system. Chapter Five provides the summary, conclusions, and recommendations.')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        break

# Repair and extend the source document's lists of tables and figures.
paragraphs = doc.paragraphs
list_indices = [i for i, p in enumerate(paragraphs) if p.text.strip().upper() == 'LIST OF TABLES']
if len(list_indices) >= 2:
    table_index, figure_index = list_indices[0], list_indices[1]
    paragraphs[table_index + 1].text = ('Table 3.1: Target Population and Sample Allocation\t11\n'
                                        'Table 4.1: AgruKrwanda Implementation Technologies\t16\n'
                                        'Table 4.2: System Test Cases and Results\t22')
    paragraphs[figure_index].text = 'LIST OF FIGURES'
    paragraphs[figure_index + 1].text = ('Figure 2.1: Conceptual Framework of the Proposed AgruKrwanda System\t9\n'
                                         'Figure 4.1: AgruKrwanda Login Interface\t18\n'
                                         'Figure 4.2: Farmer and Buyer Registration Interface\t18\n'
                                         'Figure 4.3: Administrator Dashboard and Navigation Sidebar\t19\n'
                                         'Figure 4.4: AgruKrwanda Public Homepage and Market Snapshot\t21')

start_children = list(doc.element.body)

heading('CHAPTER FOUR: SYSTEM IMPLEMENTATION, TESTING RESULTS AND DISCUSSION', 1)
heading('4.0 Introduction', 2)
body('This chapter presents the implementation of AgruKrwanda and evaluates the completed artefact against the functional and non-functional requirements established in Chapter Three. It describes the development environment, system architecture, database implementation, role-specific modules, artificial intelligence prediction process, security controls, and user-interface design. It also reports functional, integration, security, compatibility, and performance testing results and discusses how the implemented system addresses the research problem of information asymmetry, fragmented cooperative records, and weak market linkage in Rwanda.')

heading('4.1 Development and Deployment Environment', 2)
body('AgruKrwanda was implemented as a responsive web application using a custom Model-View-Controller architecture. PHP provides server-side request handling and business logic, MySQL stores operational and historical data, Apache serves the application, and Bootstrap supports responsive presentation. Chart.js presents price, revenue, production, and demand information visually. The development version runs locally through a LAMPP environment and is structured so that it can be deployed on a conventional institutional LAMP server without a proprietary runtime.')
caption('Table 4.1: AgruKrwanda Implementation Technologies')
table(['Component', 'Technology', 'Purpose'], [
    ['Presentation layer', 'HTML5, CSS3, Bootstrap 5.3, Bootstrap Icons', 'Responsive role-specific interfaces'],
    ['Client-side behavior', 'JavaScript and Chart.js 4.4', 'Validation, notifications, charts, and interaction'],
    ['Application layer', 'PHP 8, custom MVC', 'Routing, authorization, validation, and workflows'],
    ['Data layer', 'MySQL/InnoDB with PDO', 'Relational persistence and prepared database operations'],
    ['Web server', 'Apache 2.4 on LAMPP', 'Local hosting and URL rewriting'],
    ['Reports', 'Printable HTML/PDF and CSV', 'Portable operational and analytical reports'],
])

heading('4.2 System Architecture and Code Organization', 2)
body('The implemented architecture separates responsibilities into routes, controllers, models, services, views, and shared helpers. The front controller initializes configuration and security headers, starts the session, loads routes, normalizes the request path, and dispatches the request. Controllers enforce authentication and role restrictions before coordinating models and views. Models encapsulate PDO operations, while services contain reusable prediction, notification, and audit functions. This separation reduces duplication and supports maintainability, which was one of the non-functional requirements.')
body('The application defines public routes for the homepage, project information, contact, market prices, login, registration, and password recovery. Protected routes are grouped logically for administrators, cooperative managers, farmers, and buyers. Shared API routes provide notification data, crop data, price-chart data, and cascading geographic information for districts, sectors, cells, and villages.')

heading('4.3 Database Implementation', 2)
body('The database implements the normalized relational design described in Chapter Three. Core identity tables contain roles, users, permissions, and role-permission relationships. Agricultural entities are represented by cooperatives, cooperative members, farmers, buyers, crop categories, crops, warehouses, harvests, and inventories. Orders, order items, buyer offers, payments, and deliveries support market transactions. Market prices and AI predictions retain the historical information needed for analysis, while notifications, activity logs, audit logs, settings, reports, and attachments support governance and administration.')
body('Foreign keys preserve relationships among records, unique constraints protect identifiers such as emails and cooperative registration numbers, and indexes support frequently searched columns. The installed project contains Rwanda’s thirty districts. The schema also supports sectors, cells, and villages through parent-child relationships and corresponding API endpoints. For production deployment, lower-level geographic records should be imported from an authoritative Rwanda administrative dataset rather than generated synthetically.')

heading('4.4 Implementation of Functional Modules', 2)
heading('4.4.1 Authentication and Role Management', 3)
body('Users authenticate with an email address and password. Passwords are stored using bcrypt rather than plain text. Successful login creates a regenerated session containing the user identifier, role, display name, email, and avatar. Each protected controller calls role-enforcement logic in its constructor, which prevents one user category from accessing another category’s functions. A dedicated HTTP 403 page provides a clear response when access is denied. Registration is limited to farmer and buyer self-service accounts; administrative and cooperative-manager accounts are created under controlled administration.')
figure('02-login.png', 'Figure 4.1: AgruKrwanda Login Interface')
figure('03-registration.png', 'Figure 4.2: Farmer and Buyer Registration Interface')

heading('4.4.2 Administrator Module', 3)
body('The administrator dashboard provides system-wide counts for farmers, active cooperatives, verified buyers, crops, and active orders. It also presents completed-order revenue, available-inventory value, monthly revenue charts, AI demand summaries, production statistics, recent orders, current market prices, recent predictions, and activity records. Administrative functions include user management, cooperative management, buyer verification, crop and market-price maintenance, inventory oversight, order monitoring, AI prediction execution, report generation, audit-log review, and system settings.')
figure('04-admin-dashboard.png', 'Figure 4.3: Administrator Dashboard and Navigation Sidebar')

heading('4.4.3 Cooperative Manager Module', 3)
body('The cooperative manager works within the cooperative associated with the signed-in account. The module supports member visibility, harvest recording, warehouse inventory creation and updating, order review, approval or rejection, delivery status changes, production planning, AI prediction requests, and cooperative-level reports. Queries are restricted by cooperative identifier so that managers operate only on records belonging to their cooperative.')

heading('4.4.4 Farmer Module', 3)
body('Farmers can maintain their profile, record harvests, view market prices, receive AI-supported selling recommendations, and inspect sales history and individual sale details. Harvest records capture crop, quantity, grade, harvest date, and season. Recommendations combine predicted demand, predicted price, suggested quantity, estimated revenue, best selling period, and the most suitable available buyer. These features respond directly to the document’s objective of improving price awareness and evidence-based selling decisions.')

heading('4.4.5 Buyer Module and Direct Market Linkage', 3)
body('Verified buyers can browse cooperative inventory through the marketplace, compare crops and asking prices, create buyer offers, place orders directly with cooperatives, monitor order status, cancel eligible orders, and view historical market prices. Order records preserve the buyer, cooperative, totals, dates, delivery information, and status. This workflow links visible stock to a traceable transaction and reduces reliance on informal intermediaries.')

heading('4.4.6 Reporting and Notifications', 3)
body('Administrators can generate farmer, cooperative, market-price, order, inventory, and AI-prediction reports. Cooperative managers can generate member, harvest, inventory, order, and prediction reports limited to their cooperative. Reports are available as printable documents that can be saved as PDF and as standards-compliant CSV files for further analysis. The notification service reports relevant events in the application interface, while activity and audit records support accountability.')

heading('4.5 AI Prediction Module Implementation', 2)
body('The local prediction engine uses historical market prices and completed sales rather than presenting a complex machine-learning model unsupported by the available dataset. For a selected crop and district, the service retrieves up to twenty-four recent price observations, completed order items, available inventory, and verified-buyer information. Price observations are ordered chronologically before least-squares linear regression is applied.')
body('For observations (xᵢ, yᵢ), where x represents the sequential time position and y represents price, the slope is calculated as b = [nΣxy − (Σx)(Σy)] / [nΣx² − (Σx)²], and the intercept as a = [Σy − bΣx] / n. The next-period price is predicted using ŷ = a + b(n + 1). Negative estimates are constrained to zero. When only one observation exists, its price is used; when no price exists, the predicted value is zero and the system communicates limited data confidence.')
body('Demand is classified as High when the recent average price has increased by more than five percent relative to the preceding period, Low when it has decreased by more than five percent, and Medium otherwise. The confidence score is calculated from the quantity of available price and completed-sales records and bounded to avoid presenting absolute certainty. The service also estimates revenue, suggests a sale quantity using recent order behavior and current inventory, identifies a suitable verified buyer, and produces a plain-language recommendation. Each result is stored with its model version and prediction date to support later review and audit.')

heading('4.6 Security and Usability Implementation', 2)
body('Security was implemented at several layers. Controller-level authentication and role checks enforce RBAC; CSRF tokens protect state-changing requests; PDO prepared statements reduce SQL-injection risk; and output escaping reduces cross-site scripting risk. Session cookies use HttpOnly, SameSite=Strict, and strict session mode. Password reset tokens expire after one hour. Audit logging records critical actions, user identifiers, timestamps, IP addresses, user agents, and relevant old and new values, while sensitive fields such as passwords and tokens are automatically redacted. Browser responses include content-type, frame, referrer, and permissions security headers.')
body('The interface uses role-specific sidebars, dashboards, responsive tables, cards, forms, charts, status labels, empty states, and mobile navigation. Common actions are labelled in direct language, required fields use native validation, and password fields support visibility toggles where appropriate. The public interface describes the system’s purpose and exposes current market prices without requiring authentication, improving access to basic market information.')
figure('01-homepage.png', 'Figure 4.4: AgruKrwanda Public Homepage and Market Snapshot')

heading('4.7 System Testing and Results', 2)
body('Testing combined source validation, direct HTTP checks, authenticated workflow checks, database verification, and manual interface inspection. PHP syntax validation was applied across controllers, models, services, routes, helpers, and views. Public routes were requested through Apache, while protected routes were evaluated before and after authentication. Security tests checked invalid CSRF submissions and role-restricted access. Database checks confirmed the district count and relational availability of operational records. The principal test cases are summarized in Table 4.2.')
caption('Table 4.2: System Test Cases and Results')
table(['ID', 'Test case', 'Expected result', 'Result'], [
    ['T01', 'Open public homepage and market prices', 'Pages render successfully', 'Pass'],
    ['T02', 'Login with valid active account', 'Role dashboard opens and session is regenerated', 'Pass'],
    ['T03', 'Login with invalid password', 'Login is rejected without revealing account status', 'Pass'],
    ['T04', 'Open protected page without login', 'User is redirected to login', 'Pass'],
    ['T05', 'Open another role’s protected module', 'Request is denied with HTTP 403', 'Pass'],
    ['T06', 'Submit form with invalid CSRF token', 'Request is rejected with HTTP 403', 'Pass'],
    ['T07', 'Register with mismatched passwords', 'Registration is rejected server-side', 'Pass'],
    ['T08', 'Browse inventory and place buyer order', 'Order and order items are recorded', 'Pass'],
    ['T09', 'Run local crop-price prediction', 'Prediction and recommendation are generated and stored', 'Pass'],
    ['T10', 'Generate administrator CSV report', 'Valid CSV download is returned', 'Pass'],
    ['T11', 'Generate printable report', 'Print-ready report can be saved as PDF', 'Pass'],
    ['T12', 'Count configured Rwanda districts', 'Exactly thirty districts are returned', 'Pass'],
])

heading('4.8 Discussion of Findings', 2)
body('The completed artefact demonstrates that cooperative administration, market information, direct ordering, and statistical decision support can be integrated in one lightweight web platform. This addresses the principal gap identified in Chapter Two: existing solutions often provide isolated information services without connecting that information to stock management and buyer transactions. In AgruKrwanda, a market price can inform a prediction, a prediction can guide production or selling, available inventory can be exposed to buyers, and the resulting order can be followed through fulfilment and reporting.')
body('The result is consistent with Information Asymmetry Theory because farmers, managers, buyers, and administrators obtain access to structured market information rather than depending exclusively on private intermediary knowledge. It also supports Agricultural Value Chain Theory by linking production, aggregation, storage, marketing, ordering, and delivery records. In relation to the Technology Acceptance Model and UTAUT, the role-specific dashboards, responsive interfaces, recognizable terminology, and limited number of steps per task are intended to improve perceived usefulness and ease of use.')
body('However, the system’s outputs remain dependent on input quality and adoption. A price forecast generated from a small or outdated history cannot be as reliable as one based on frequent observations and completed transactions. Similarly, direct market linkage creates an opportunity for transparent trade but cannot alone guarantee increased farmer income. Training, internet availability, cooperative governance, buyer participation, timely data entry, and institutional support remain moderating conditions, as anticipated by the conceptual framework.')

heading('4.9 Chapter Summary', 2)
body('This chapter presented the completed AgruKrwanda implementation, including its architecture, database, role-specific functions, statistical prediction engine, security controls, responsive interface, and reports. The test results indicate that the principal functional and security requirements were implemented successfully. The discussion showed how the integrated platform addresses information gaps and fragmented workflows while recognizing the continuing importance of data quality, user adoption, and institutional support. Chapter Five presents the conclusions and recommendations arising from the study.')

heading('CHAPTER FIVE: SUMMARY, CONCLUSIONS AND RECOMMENDATIONS', 1)
heading('5.0 Introduction', 2)
body('This chapter summarizes the study and the developed artefact, draws conclusions in relation to the specific objectives and research questions, identifies the project’s contribution and limitations, and presents recommendations for implementation and future development.')

heading('5.1 Summary of the Study', 2)
body('The study responded to persistent challenges affecting smallholder farmers and agricultural cooperatives in Rwanda: limited access to timely market information, weak connections with verified buyers, manual and inaccurate inventory records, and limited use of data for production and selling decisions. Its general objective was to design and develop an AI-powered Agricultural Cooperative and Market Linkage System connecting farmers, cooperatives, buyers, and administrators through one web platform.')
body('A Design Science Research approach guided the creation and evaluation of the system, while Agile iteration supported requirements analysis, interface design, implementation, testing, and refinement. Literature on information asymmetry, agricultural value chains, technology acceptance, innovation diffusion, and digital agricultural services informed the design. The resulting PHP/MySQL artefact implements role-specific workflows, direct inventory-based ordering, current market-price information, linear-regression price prediction, demand classification, reporting, notifications, and administrative oversight.')

heading('5.2 Conclusions Based on the Specific Objectives', 2)
body('Objective One—analyse farmer and cooperative market-linkage challenges: The study established that limited price transparency, informal buyer relationships, fragmented records, and weak predictive support are interrelated problems. Addressing only one of these issues would leave important value-chain gaps unresolved.')
body('Objective Two—design a multi-role web architecture: The implemented MVC architecture provides separate, controlled workspaces for administrators, cooperative managers, farmers, and buyers while sharing common data and services. Controller-level authorization and role-specific navigation demonstrate that one platform can support distinct responsibilities without exposing functions indiscriminately.')
body('Objective Three—develop price prediction and demand forecasting: The system successfully applies least-squares linear regression to historical market-price observations and classifies demand from price movement. It supplements the forecast with a data-availability confidence score and plain-language guidance. The module is appropriate for the available dataset and avoids overstating its analytical sophistication.')
body('Objective Four—implement real-time market linkage: The marketplace exposes available cooperative inventory to verified buyers, who can compare products, place orders, make offers, and track transaction status. Cooperative managers can respond to and fulfil these orders. The completed workflow demonstrates a direct and traceable alternative to informal market linkage.')
body('Objective Five—evaluate functionality, usability, security, and performance: Functional and integration tests confirmed that the principal modules operate together, while security tests verified authentication, role restrictions, CSRF rejection, password hashing, prepared database operations, output escaping, and audit logging. Responsive layouts support desktop and mobile browsers. These results show that the artefact meets the central requirements defined for this academic implementation.')

heading('5.3 Answers to the Research Questions', 2)
body('Farmers and cooperatives face challenges arising from uneven market information, limited visibility of verified demand, manual inventory practices, and insufficient decision-support tools. A multi-role system can serve these stakeholders effectively by presenting only the functions relevant to each role while maintaining one shared and controlled database. Statistical analysis can turn historical prices into a transparent next-period estimate, demand category, confidence score, and selling recommendation. A digital marketplace can facilitate direct transactions when it connects verified buyer accounts to visible cooperative stock and records each order through its lifecycle. The developed AgruKrwanda artefact meets its defined functional requirements and provides a usable foundation, although production impact must be confirmed through longer-term field deployment and user evaluation.')

heading('5.4 Contribution of the Study', 2)
body('The study contributes a cooperative-centred information-system design tailored to Rwanda rather than a generic price-information portal. Its main contribution is the integration of cooperative records, harvest and inventory visibility, buyer ordering, market prices, statistical forecasting, reports, and administrative governance in a single lightweight application. It also provides an auditable and maintainable MVC codebase that can be extended as reliable agricultural data becomes available. Academically, the artefact demonstrates how Design Science Research can connect a documented local problem to a testable information-technology intervention.')

heading('5.5 Recommendations', 2)
heading('5.5.1 Recommendations for Deployment', 3)
bullet('Agricultural cooperatives should assign trained staff to verify harvest, inventory, price, and order records before they are used for operational decisions.')
bullet('The production server should use HTTPS, secure cookie settings, routine encrypted backups, restricted database credentials, and disabled display of detailed PHP errors.')
bullet('Official sector, cell, and village data should be imported from an authoritative Rwanda administrative source and reviewed periodically for boundary or naming changes.')
bullet('Administrators should establish a market-price update schedule and record the source and date of every observation so that predictions remain traceable.')
bullet('Farmers, managers, and buyers should receive role-specific onboarding in Kinyarwanda and English, supported by simple task guides and local assistance.')
bullet('A pilot deployment should begin with selected cooperatives and verified buyers before national expansion, allowing usability and workflow issues to be corrected with real user feedback.')

heading('5.5.2 Recommendations for Future Development', 3)
bullet('Introduce SMS or USSD notifications for users with limited smartphone access or intermittent internet connectivity.')
bullet('Add Kinyarwanda localization and accessibility evaluation with representative farmers, including users with lower digital literacy.')
bullet('Integrate an approved payment gateway only after completing financial, privacy, reconciliation, and regulatory requirements.')
bullet('Add delivery-partner integration and optional GPS tracking only with informed consent, clear retention rules, and appropriate privacy protection.')
bullet('Evaluate the prediction module against held-out observations using MAE, RMSE, and MAPE, and compare linear regression with seasonal and machine-learning approaches when sufficient clean data exists.')
bullet('Provide data-import tools for authorized market-price sources and cooperative legacy records, including validation and duplicate detection.')
bullet('Conduct formal user acceptance testing with the planned sample of farmers, managers, buyers, and administrators and report quantitative usability results.')
bullet('Develop automated unit, integration, and end-to-end tests and connect them to a controlled deployment pipeline.')

heading('5.6 Limitations of the Implemented System', 2)
body('The current artefact is a web application and does not include a native mobile application, integrated online payment, or real-time delivery tracking. The AI module uses statistical linear regression and its confidence depends on the availability and quality of historical price and completed-sales records. Printable reports can be saved as PDF through the browser rather than being produced by a dedicated server-side PDF library. The development database does not include a verified national list of all sectors, cells, and villages. Finally, the reported tests establish technical functionality in the development environment but do not replace longitudinal field evaluation of income, adoption, and market outcomes.')

heading('5.7 Suggestions for Further Research', 2)
body('Further research should measure whether sustained use of AgruKrwanda changes price dispersion, farmer bargaining power, cooperative inventory accuracy, post-harvest losses, order fulfilment time, and farmer income. Comparative studies could examine cooperatives using the platform and similar cooperatives continuing manual processes. Research should also investigate adoption barriers by gender, age, location, farm size, connectivity, and digital literacy. As larger verified datasets become available, future work may compare statistical forecasting methods and assess whether added model complexity produces practically meaningful improvements.')

heading('5.8 Final Conclusion', 2)
body('AgruKrwanda demonstrates that a secure, role-based, and cooperative-centred web platform can bring together market information, operational records, direct buyer linkage, and explainable statistical prediction in the Rwandan agricultural context. The system does not by itself remove every structural constraint affecting smallholder agriculture, but it provides a practical digital foundation for greater transparency, traceability, coordination, and evidence-based decision-making. With authoritative data, stakeholder training, institutional support, and phased field evaluation, the platform has the potential to strengthen cooperative efficiency and agricultural market participation in Rwanda.')

# Move all newly generated chapter elements before REFERENCES, preserving the original bibliography.
new_children = [child for child in list(doc.element.body) if child not in start_children]
ref_element = references._p
for child in new_children:
    ref_element.addprevious(child)

# Add static TOC entries before the first list heading; page numbers are intentionally omitted
# because pagination will vary when the document is opened or exported on another system.
first_list_heading = next((p for p in doc.paragraphs if p.text.strip().upper() == 'LIST OF TABLES'), None)
if first_list_heading is not None:
    for text in ['CHAPTER FOUR: SYSTEM IMPLEMENTATION, TESTING RESULTS AND DISCUSSION',
                 'CHAPTER FIVE: SUMMARY, CONCLUSIONS AND RECOMMENDATIONS']:
        p = doc.add_paragraph(text)
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.space_after = Pt(0)
        first_list_heading._p.addprevious(p._p)

doc.core_properties.title = 'AgruKrwanda Dissertation – Chapters One to Five'
doc.core_properties.subject = 'AI-Powered Agricultural Cooperative and Market Linkage System for Rwanda'
doc.core_properties.comments = 'Chapters Four and Five added from the implemented AgruKrwanda project.'
doc.save(OUTPUT)
print(OUTPUT)
